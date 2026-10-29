"""Self-test for `file_size` over a throwaway git repository."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfWriter

from pystatic import config, files
from pystatic.harness.docref import doc_anchor_exists
from pystatic.scripts import file_size as rule

GIT = shutil.which("git") or "git"


def _repo(tmp_path: Path, sizes: dict[str, int], table: str) -> None:
    for rel, size in sizes.items():
        path = tmp_path / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"x" * size)
    subprocess.run([GIT, "init", "-q"], cwd=tmp_path, check=True)
    (tmp_path / ".gitignore").write_text("ignored/\n")
    cfg = tmp_path / "pystatic.toml"
    cfg.write_text(table)
    files.configure(tmp_path)
    config.configure(cfg)


def test_doc_anchor_resolves() -> None:
    assert doc_anchor_exists(rule.DOC) is True


def test_only_watched_tracked_files_over_the_ceiling_are_reported(tmp_path: Path) -> None:
    _repo(
        tmp_path,
        {"data/big.json": 101, "data/small.json": 100, "ignored/big.bin": 500, "src/big.py": 500},
        '[file-size]\npaths = ["data", "ignored"]\nmax-bytes = 100\n',
    )
    findings = rule.check()
    assert [(f.path, f.rule) for f in findings] == [("data/big.json", "file-size-over-ceiling")]


def test_a_bad_ceiling_is_a_configuration_error(tmp_path: Path) -> None:
    _repo(tmp_path, {}, "[file-size]\nmax-bytes = 0\n")
    with pytest.raises(config.ConfigError, match="max-bytes"):
        rule.check()


def _pdf(path: Path, pages: int) -> None:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("wb") as output:
        writer.write(output)


def _deck(path: Path, slides: int) -> None:
    deck = Presentation()
    for _ in range(slides):
        deck.slides.add_slide(deck.slide_layouts[6])
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))


def test_pages_over_the_budget_and_uncountable_documents_are_reported(tmp_path: Path) -> None:
    _pdf(tmp_path / "corpus/ok.pdf", 2)
    _pdf(tmp_path / "corpus/long.pdf", 3)
    _deck(tmp_path / "corpus/long.pptx", 3)
    _repo(
        tmp_path,
        {"corpus/notes.docx": 8},
        '[file-size]\npaths = []\npage-paths = ["corpus"]\nmax-pages = 2\n',
    )
    findings = rule.check()
    assert [(f.path, f.rule) for f in findings] == [
        ("corpus/long.pdf", "file-size-over-page-budget"),
        ("corpus/long.pptx", "file-size-over-page-budget"),
        ("corpus/notes.docx", "file-size-pages-not-countable"),
    ]


def test_suffixes_limit_the_judged_files(tmp_path: Path) -> None:
    _repo(
        tmp_path,
        {"data/big.png": 200, "data/big.json": 200},
        '[file-size]\npaths = ["data"]\nmax-bytes = 100\nsuffixes = [".png"]\n',
    )
    assert [f.path for f in rule.check()] == ["data/big.png"]
